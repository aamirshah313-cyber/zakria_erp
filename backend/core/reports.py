import csv
import io
from collections import defaultdict
from decimal import Decimal
from xml.sax.saxutils import escape
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.exceptions import ValidationError
from .models import Document, Company, SavedReport
from .auth import require, audit
from .serializers import CompanySerializer, PartySerializer

FIELDS = {'number': 'Document number', 'party': 'Customer', 'issue_date': 'Issue date', 'due_date': 'Due / valid until', 'status': 'Status', 'reference': 'Reference', 'subtotal': 'Subtotal', 'tax_total': 'Tax', 'total': 'Total (PKR)', 'owner': 'Prepared by'}

def safe_cell(value):
    if isinstance(value, (int, float, Decimal)):
        return value
    text = str(value if value is not None else '')
    return "'" + text if text.startswith(('=', '+', '-', '@', '\t', '\r')) else text

def csv_response(rows, fields, filename):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(fields)
    for row in rows:
        writer.writerow([safe_cell(row.get(f, '')) for f in fields])
    response = HttpResponse('\ufeff' + out.getvalue(), content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response

def report_rows(user, definition):
    kind = definition.get('kind', 'quotation')
    if kind not in ['quotation', 'invoice']:
        raise ValidationError('Choose a valid dataset.')
    require(user, f'{kind}.view')
    columns = definition.get('columns', list(FIELDS))
    if not isinstance(columns, list) or not columns or any(c not in FIELDS for c in columns):
        raise ValidationError('Choose valid report columns.')
    query = Document.objects.filter(kind=kind).exclude(status='deleted').select_related('party', 'owner').order_by('-issue_date', '-id')
    for key, field in [('from', 'issue_date__gte'), ('to', 'issue_date__lte')]:
        if definition.get(key):
            from datetime import date
            try:
                date.fromisoformat(definition[key])
            except (ValueError, TypeError):
                raise ValidationError('Dates must use YYYY-MM-DD.')
            query = query.filter(**{field: definition[key]})
    if definition.get('status'):
        query = query.filter(status=definition['status'])
    rows = []
    if query.count() > 5000:
        raise ValidationError('Narrow the date range to 5,000 documents or fewer for this pilot.')
    for d in query:
        rows.append({'number': d.number, 'party': d.issued_snapshot.get('party', {}).get('name', d.party.name), 'issue_date': str(d.issue_date), 'due_date': str(d.due_date or ''), 'status': d.status, 'reference': d.reference, 'subtotal': d.subtotal, 'tax_total': d.tax_total, 'total': d.total, 'owner': d.owner.username})
    group = definition.get('group', 'status')
    if group not in ['status', 'party', 'month']:
        raise ValidationError('Group by status, customer or month.')
    totals = defaultdict(Decimal)
    for row in rows:
        key = row['issue_date'][:7] if group == 'month' else row[group]
        totals[key] += row['total']
    return columns, rows, [{'label': key, 'value': value} for key, value in sorted(totals.items())]

def pdf_table(title, headers, rows, options, footer=''):
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, A3, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    import re
    accent = options.get('accent', '#176B61')
    if not re.fullmatch(r'#[0-9a-fA-F]{6}', accent):
        raise ValidationError('Use a six-digit hex colour.')
    size = A3 if options.get('paper') == 'A3' else A4
    if options.get('orientation') == 'landscape':
        size = landscape(size)
    buf = io.BytesIO()
    from .branding import PDF_HEADER_SPACE, draw_pdf_logo
    doc = SimpleDocTemplate(buf, pagesize=size, rightMargin=32, leftMargin=32, topMargin=PDF_HEADER_SPACE + 16, bottomMargin=45)
    styles = getSampleStyleSheet()
    styles['BodyText'].fontSize = 8
    styles['BodyText'].leading = 11
    p = lambda s: Paragraph(escape(str(s)).replace('\n', '<br/>'), styles['BodyText'])
    story = [Paragraph(escape(title), styles['Title']), Spacer(1, 18)]
    if options.get('intro'):
        story += [p(options['intro']), Spacer(1, 14)]
    table = Table([[p(h) for h in headers]] + [[p(v) for v in row] for row in rows], repeatRows=1, colWidths=[(size[0] - 64) / len(headers)] * len(headers))
    table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(accent)), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F2F5F5')]), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8), ('LINEBELOW', (0, 0), (-1, 0), .5, colors.grey)]))
    # Header text uses a light background for readable black text in print.
    table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DDECE9'))]))
    story.append(table)
    if options.get('outro'):
        story += [Spacer(1, 16), p(options['outro'])]
    if footer:
        story += [Spacer(1, 16), p(footer)]
    def page(canvas, document):
        draw_pdf_logo(canvas, size[0], size[1], 32, colors.HexColor(accent))
        canvas.setFont('Helvetica', 8)
        if options.get('draft'):
            canvas.drawString(32, 22, 'DRAFT - NOT ISSUED')
        canvas.drawRightString(size[0] - 32, 22, f'Page {document.page}')
    doc.build(story, onFirstPage=page, onLaterPages=page)
    return buf.getvalue()

class Reports(APIView):
    def get(self, request):
        require(request.user, 'reports.view')
        return Response({'fields': FIELDS, 'saved': list(SavedReport.objects.filter(owner=request.user).values('id', 'name', 'definition'))})
    def post(self, request):
        require(request.user, 'reports.view')
        definition = request.data
        columns, rows, chart = report_rows(request.user, definition)
        if definition.get('save_name'):
            name = str(definition['save_name']).strip()[:100]
            SavedReport.objects.create(owner=request.user, name=name, definition={k: v for k, v in definition.items() if k not in ['save_name', 'format']})
        output = definition.get('format', 'json')
        if output == 'json':
            return Response({'columns': columns, 'rows': rows, 'chart': chart, 'count': len(rows), 'total': sum((r['total'] for r in rows), Decimal(0))})
        require(request.user, 'reports.export')
        if output == 'csv':
            response = csv_response(rows, columns, 'report.csv')
        elif output == 'xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from .branding import add_xlsx_banner
            wb = Workbook()
            ws = wb.active
            ws.title = 'Report'
            first = add_xlsx_banner(ws)
            ws.append([FIELDS[c] for c in columns])
            for row in rows:
                ws.append([safe_cell(row[c]) for c in columns])
            from openpyxl.utils import get_column_letter
            ws.freeze_panes = f'A{first + 1}'
            ws.auto_filter.ref = f'A{first}:{get_column_letter(len(columns))}{ws.max_row}'
            for cell in ws[first]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='176B61')
            for i, c in enumerate(columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = 24
                if c in ['subtotal', 'tax_total', 'total']:
                    for row in ws.iter_rows(min_row=first + 1, min_col=i, max_col=i):
                        row[0].number_format = '#,##0.00'
            stream = io.BytesIO()
            wb.save(stream)
            response = HttpResponse(stream.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = 'attachment; filename="report.xlsx"'
        elif output == 'pdf':
            response = HttpResponse(pdf_table(f"{definition.get('kind', 'quotation').title()} report", [FIELDS[c] for c in columns], [[r[c] for c in columns] for r in rows], definition, 'Document totals are not a financial statement.'), content_type='application/pdf')
            response['Content-Disposition'] = 'attachment; filename="report.pdf"'
        else:
            raise ValidationError('Supported export formats: PDF, XLSX and CSV.')
        audit(request.user, 'report.exported', definition.get('kind'), format=output, count=len(rows))
        return response

class DocumentPDF(APIView):
    def get(self, request, pk):
        doc = get_object_or_404(Document.objects.select_related('party'), pk=pk)
        require(request.user, f'{doc.kind}.view')
        require(request.user, 'reports.export')
        snapshot = doc.issued_snapshot
        company = snapshot.get('company', CompanySerializer(Company.objects.get(pk=1)).data)
        party = snapshot.get('party', PartySerializer(doc.party).data)
        title = f"{doc.kind.upper()} | {doc.number}"
        if doc.status != 'issued':
            title = 'DRAFT - NOT ISSUED\n' + title
        intro = f"{company['name']}\n{company['address']}\n{company['city']}\n"
        intro += ' | '.join(f'{k.upper()}: {company[k]}' for k in ['ntn', 'strn', 'ftn'] if company.get(k))
        intro += f"\n\nCustomer: {party['name']}\n{party['address']}\nIssue date: {doc.issue_date} | Due / valid until: {doc.due_date or '-'}\nReference: {doc.reference or '-'}\nTax treatment: {doc.tax_treatment}"
        rows = [[l.description, l.unit, str(l.quantity), f'{l.rate:,.2f}', f'{l.tax_amount:,.2f}', f'{l.amount + l.tax_amount:,.2f}'] for l in doc.lines.all()]
        outro = f"Subtotal: PKR {doc.subtotal:,.2f}\nTax: PKR {doc.tax_total:,.2f}\nTOTAL: PKR {doc.total:,.2f}\n\n{doc.notes}\n\n{company['bank_details']}"
        options = dict(request.query_params.items())
        options.update(intro=intro, outro=outro, accent=company['accent'], draft=doc.status != 'issued')
        data = pdf_table(title, ['Description', 'Unit', 'Quantity', 'Unit price', 'Tax', 'Amount'], rows, options, company['footer'])
        response = HttpResponse(data, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{doc.number}.pdf"'
        audit(request.user, 'document.exported', doc.number, status=doc.status)
        return response
