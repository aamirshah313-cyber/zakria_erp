"""Bounded spreadsheet staging. No workbook formulas/macros or approvals execute."""
import csv
import hashlib
import io
import re
import uuid
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath
from django.db import transaction
from django.shortcuts import get_object_or_404
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.response import Response
from rest_framework.views import APIView
from .auth import require, audit
from .models import Company, RegisterEntry, RegisterCategory, RegisterSource, Project, Party, RegisterImport, RegisterImportOrigin
from .recovery import desktop_only
from .registers import EntrySerializer

MAX_ROWS, MAX_COLUMNS = 1000, 40
FIELDS = ['date', 'party', 'amount', 'direction', 'payment', 'receipt', 'category', 'source', 'project', 'method', 'reference', 'handled_by', 'remarks', 'beneficiary', 'nature', 'counterparty', 'reporting_class']


def access(request):
    desktop_only()
    for permission in ['register.view', 'register.create', 'register.import']:
        require(request.user, permission)


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    result = str(value)
    if len(result) > 4000:
        raise ValidationError('A cell exceeds 4,000 characters; shorten it before import.')
    return result.strip()


def read_upload(request):
    upload = request.FILES.get('file')
    if not upload or not 0 < upload.size <= 5 * 1024 * 1024:
        raise ValidationError('Select an XLSX or UTF-8 CSV file up to 5 MB.')
    content = upload.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise ValidationError('File exceeds 5 MB.')
    name = PurePosixPath(upload.name.replace('\\', '/')).name
    if len(name) > 180:
        raise ValidationError('Filename is too long.')
    extension = PurePosixPath(name).suffix.lower()
    sheets, warnings = {}, []
    try:
        if extension == '.csv':
            data = list(csv.reader(io.StringIO(content.decode('utf-8-sig'))))
            if len(data) > MAX_ROWS + 50 or any(len(row) > MAX_COLUMNS for row in data):
                raise ValidationError('Maximum 1,050 source rows and 40 columns per file.')
            sheets['CSV'] = [[cell_text(cell) for cell in row] for row in data]
        elif extension == '.xlsx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                members = archive.infolist()
                if len(members) > 2000 or sum(item.file_size for item in members) > 32 * 1024 * 1024:
                    raise ValidationError('Workbook exceeds the safe expanded-size limit (32 MB).')
                if any('vbaproject' in item.filename.lower() for item in members):
                    raise ValidationError('Macro content is not supported.')
                for item in members:
                    if item.filename.endswith('.xml'):
                        xml = archive.read(item).upper()
                        if b'<!DOCTYPE' in xml or b'<!ENTITY' in xml:
                            raise ValidationError('XML document types/entities are not supported.')
                if any(item.filename.startswith(('xl/media/', 'xl/embeddings/')) for item in members):
                    warnings.append('Embedded images/objects are not imported. Upload supporting documents to the resulting drafts separately.')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) > 20:
                    raise ValidationError('Select a workbook with no more than 20 sheets.')
                for sheet in book.worksheets:
                    if (sheet.max_row is not None and sheet.max_row > MAX_ROWS + 50) or (sheet.max_column is not None and sheet.max_column > MAX_COLUMNS):
                        raise ValidationError('Maximum 1,050 source rows and 40 columns per sheet; remove unused formatted rows/columns.')
                    # Dimensions are optional in XLSX and may be stale. Stream
                    # actual rows with limits instead of trusting that metadata.
                    sheet.reset_dimensions()
                    values = []
                    for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if index > MAX_ROWS + 50 or len(row) > MAX_COLUMNS:
                            raise ValidationError('Maximum 1,050 source rows and 40 columns per sheet.')
                        values.append([cell_text(cell) for cell in row])
                    sheets[sheet.title] = values
            finally:
                book.close()
        else:
            raise ValidationError('Use XLSX or UTF-8 CSV. Convert legacy XLS files first.')
    except ValidationError:
        raise
    except Exception:
        raise ValidationError('The spreadsheet could not be read. Use an unencrypted XLSX or UTF-8 CSV file.')
    if not sheets or not any(sheets.values()):
        raise ValidationError('The spreadsheet contains no rows.')
    return name, hashlib.sha256(content).hexdigest(), sheets, warnings


def exact_amount(value, blank_zero=False):
    text = str(value).strip()
    if not text and blank_zero:
        return Decimal('0.00')
    # Accept plain decimal or conventional comma grouping, never silently round.
    if not re.fullmatch(r'(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d{1,2})?', text):
        raise ValueError('Use a non-negative amount with at most two decimal places.')
    try:
        amount = Decimal(text.replace(',', ''))
        if amount > Decimal('99999999999999.99'):
            raise ValueError('Amount is too large.')
        return amount
    except InvalidOperation:
        raise ValueError('Invalid amount.')


def parsed_date(value, mode):
    formats = ['%Y-%m-%d']
    if mode == 'day_first':
        formats += ['%d/%m/%Y', '%d/%m/%y', '%d-%m-%Y', '%d-%m-%y', '%d.%m.%Y', '%d.%m.%y', '%d-%b-%y', '%d-%b-%Y', '%d/%b/%y', '%d/%b/%Y']
    elif mode == 'month_first':
        formats += ['%m/%d/%Y', '%m/%d/%y', '%m-%d-%Y', '%m-%d-%y']
    for fmt in formats:
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError('Invalid date for the selected date format. Excel date cells or ISO YYYY-MM-DD are supported.')


def checked_config(config, batch):
    if not isinstance(config, dict):
        raise ValidationError('Invalid import configuration.')
    columns = config.get('columns', {})
    defaults = config.get('defaults', {})
    if not isinstance(columns, dict) or not isinstance(defaults, dict):
        raise ValidationError('Invalid column mapping.')
    for key, index in columns.items():
        if key not in FIELDS or type(index) is not int or not 0 <= index < len(batch.headers):
            raise ValidationError('A mapped column is invalid.')
    if 'date' not in columns or 'party' not in columns:
        raise ValidationError('Map Transaction date and Organization / person.')
    if config.get('amount_mode') not in ['single', 'separate'] or config.get('date_format') not in ['iso', 'day_first', 'month_first']:
        raise ValidationError('Choose amount and date formats.')
    if config['amount_mode'] == 'single' and 'amount' not in columns:
        raise ValidationError('Map Amount for a single amount column.')
    if config['amount_mode'] == 'separate' and not {'payment', 'receipt'} <= columns.keys():
        raise ValidationError('Map both Payment and Receipt columns explicitly.')
    for key in ['exclude', 'duplicate_reasons', 'aliases']:
        if not isinstance(config.get(key, {}), dict):
            raise ValidationError(f'Invalid {key}.')
    valid_rows = {str(row['row']) for row in batch.rows}
    if not set(config.get('exclude', {})) <= valid_rows or not set(config.get('duplicate_reasons', {})) <= valid_rows:
        raise ValidationError('Row decisions must refer to rows in this batch.')
    for reasons in [config.get('exclude', {}), config.get('duplicate_reasons', {})]:
        if any(not isinstance(value, str) or not 5 <= len(value.strip()) <= 500 for value in reasons.values()):
            raise ValidationError('Each exclusion or duplicate override needs a reason of 5–500 characters.')
    aliases = config.get('aliases', {})
    if any(key not in ['category', 'source', 'project', 'counterparty'] or not isinstance(value, dict) for key, value in aliases.items()):
        raise ValidationError('Invalid master-record mappings.')
    return config


def make_preview(batch, config):
    columns, defaults = config['columns'], config.get('defaults', {})
    rows, seen = [], {}
    totals = {'receipt': Decimal('0'), 'payment': Decimal('0')}
    masters = {'category': list(RegisterCategory.objects.all()), 'source': list(RegisterSource.objects.all()), 'project': list(Project.objects.all()), 'counterparty': list(Party.objects.all())}
    source_rows = set(RegisterImportOrigin.objects.filter(file_hash=batch.file_hash, sheet=batch.sheet).values_list('source_row', flat=True))
    excluded = config.get('exclude', {})
    for raw in batch.rows:
        number, values = raw['row'], raw['values']
        item = {'row': number, 'errors': [], 'duplicates': [], 'excluded': str(number) in excluded, 'reason': excluded.get(str(number), ''), 'payload': {}, 'labels': {}}
        rows.append(item)
        if item['excluded']:
            continue
        def value(key):
            index = columns.get(key)
            return values[index].strip() if index is not None and index < len(values) else ''
        try:
            if any(value(key).startswith('=') for key in columns):
                raise ValueError('A mapped cell contains a formula. Supply verified values or exclude the row as a control/summary.')
            payload = {key: value(key) for key in ['party', 'reference', 'handled_by', 'remarks', 'beneficiary']}
            payload['reporting_class'] = value('reporting_class').casefold() if 'reporting_class' in columns else defaults.get('reporting_class', 'unclassified')
            payload['nature'] = value('nature').casefold() if 'nature' in columns else defaults.get('nature', 'unclassified')
            payload['date'] = parsed_date(value('date'), config['date_format'])
            if config['amount_mode'] == 'separate':
                receipt, payment = exact_amount(value('receipt'), True), exact_amount(value('payment'), True)
                if bool(receipt) == bool(payment):
                    raise ValueError('Exactly one Receipt or Payment amount must be positive on a transaction row.')
                direction, amount = ('receipt', receipt) if receipt else ('payment', payment)
            else:
                direction = (value('direction') if 'direction' in columns else str(defaults.get('direction', ''))).casefold()
                if direction not in ['receipt', 'payment']:
                    raise ValueError('Direction must be receipt or payment; debit/credit meanings must be mapped explicitly.')
                amount = exact_amount(value('amount'))
            if amount <= 0:
                raise ValueError('Transaction amount must be positive.')
            payload.update(direction=direction, amount=f'{amount:.2f}')
            for kind, records in masters.items():
                text = value(kind)
                if kind in columns:
                    if not text and kind in ['project', 'counterparty']:
                        payload[kind] = None
                        continue
                    alias = config.get('aliases', {}).get(kind, {}).get(text)
                    candidates = [r for r in records if str(r.pk) == str(alias)] if alias is not None else [r for r in records if text and text.casefold() in [r.name.casefold(), getattr(r, 'code', '').casefold()]]
                else:
                    key = defaults.get(kind)
                    if kind in ['project', 'counterparty'] and not key:
                        payload[kind] = None
                        continue
                    candidates = [r for r in records if str(r.pk) == str(key)]
                if len(candidates) != 1 or not candidates[0].active:
                    raise ValueError(f'Map {kind} "{text}" to one active record or choose an active default.')
                payload[kind] = candidates[0].pk
                item['labels'][kind] = candidates[0].name
            payload['method'] = (value('method') if 'method' in columns else str(defaults.get('method', ''))).casefold()
            payload['request_key'] = str(uuid.uuid5(uuid.NAMESPACE_URL, f'zakaria-import:{batch.pk}:{number}'))
            validator = EntrySerializer(data=payload)
            validator.is_valid(raise_exception=True)
            item['payload'] = payload
            totals[direction] += amount
            if number in source_rows:
                item['errors'].append('This exact workbook/sheet/source row was already imported, including cancelled entries.')
            matches = RegisterEntry.objects.filter(date=payload['date'], direction=direction, amount=amount).order_by('pk')
            item['duplicates'] = [{'reference': f'REG-{r.pk}', 'party': r.party, 'status': r.status} for r in matches[:5]]
            signature = (payload['date'], direction, payload['amount'])
            if signature in seen:
                item['duplicates'].append({'reference': f'Source row {seen[signature]}', 'party': '', 'status': 'same batch'})
            seen[signature] = number
            if item['duplicates'] and not config.get('duplicate_reasons', {}).get(str(number)):
                item['errors'].append('Possible duplicate: exclude this row or record why it is a separate transaction.')
        except (ValueError, ValidationError) as exc:
            item['errors'].append(str(exc))
    errors = []
    if not any(not row['excluded'] for row in rows):
        errors.append('Include at least one transaction row.')
    for direction in ['receipt', 'payment']:
        try:
            control = exact_amount(config.get(f'control_{direction}', ''))
            if control != totals[direction]:
                errors.append(f'{direction.title()} control total differs from included valid rows by {control - totals[direction]:.2f}.')
        except ValueError:
            errors.append(f'Enter the verified {direction} control total, including 0 where applicable.')
    return {'rows': rows, 'totals': {key: f'{val:.2f}' for key, val in totals.items()}, 'errors': errors, 'ready': not errors and not any(row['errors'] for row in rows), 'included': sum(not row['excluded'] for row in rows), 'excluded': sum(row['excluded'] for row in rows)}


def detail(batch):
    return {'id': batch.pk, 'filename': batch.filename, 'file_hash': batch.file_hash, 'sheet': batch.sheet, 'header_row': batch.header_row, 'headers': batch.headers, 'rows': batch.rows, 'warnings': batch.warnings, 'configuration': batch.configuration, 'preview': batch.preview, 'version': batch.version, 'status': batch.status, 'entries': list(batch.origins.order_by('source_row').values('source_row', 'entry_id'))}


class ImportInspect(APIView):
    def post(self, request):
        access(request)
        name, digest, sheets, warnings = read_upload(request)
        return Response({'filename': name, 'sheets': [{'name': key, 'sample': value[:12]} for key, value in sheets.items()], 'warnings': warnings})


class ImportBatches(APIView):
    def get(self, request):
        access(request)
        return Response(list(RegisterImport.objects.filter(owner=request.user).exclude(status='deleted').order_by('-pk').values('id', 'filename', 'sheet', 'status', 'created_at')[:50]))

    @transaction.atomic
    def post(self, request):
        access(request)
        name, digest, sheets, warnings = read_upload(request)
        selected = request.data.get('sheet')
        try:
            header = int(request.data.get('header_row', 0))
        except (ValueError, TypeError):
            raise ValidationError('Select a valid header row.')
        if selected not in sheets or not 1 <= header <= 50 or header >= len(sheets[selected]):
            raise ValidationError('Select a sheet and header row (1–50) with transaction rows below it.')
        source = sheets[selected]
        width = max(len(row) for row in source)
        headers = source[header-1] + [''] * (width-len(source[header-1]))
        rows = [{'row': i+1, 'values': row + [''] * (width-len(row))} for i, row in enumerate(source) if i >= header and any(cell.strip() for cell in row)]
        if not rows or len(rows) > MAX_ROWS:
            raise ValidationError('Stage 1–1,000 non-empty rows at a time.')
        batch = RegisterImport.objects.create(owner=request.user, filename=name, file_hash=digest, sheet=selected, header_row=header, headers=headers, rows=rows, warnings=warnings)
        audit(request.user, 'register.import_staged', batch.pk, source_hash=digest, sheet=selected, count=len(rows))
        return Response(detail(batch), status=201)


class ImportReview(APIView):
    def get(self, request, pk):
        access(request)
        return Response(detail(get_object_or_404(RegisterImport, pk=pk, owner=request.user)))

    @transaction.atomic
    def post(self, request, pk):
        access(request)
        # Same lock as ordinary entry creation; recheck duplicates at commit.
        get_object_or_404(Company.objects.select_for_update(), pk=1)
        batch = get_object_or_404(RegisterImport.objects.select_for_update(), pk=pk, owner=request.user)
        action = request.data.get('action')
        if batch.status == 'deleted':
            raise ValidationError('Restore this removed batch through Data management before continuing.')
        if batch.status == 'imported':
            if action == 'commit':
                return Response(detail(batch))
            raise ValidationError('An imported batch is read-only.')
        if request.data.get('version') != batch.version:
            raise ValidationError('This batch changed. Reopen it before continuing.')
        if action == 'preview':
            batch.configuration = checked_config(request.data.get('configuration'), batch)
            batch.preview = make_preview(batch, batch.configuration)
            batch.status = 'reviewed'
            batch.version += 1
            batch.save()
            audit(request.user, 'register.import_reviewed', batch.pk, version=batch.version, ready=batch.preview['ready'])
        elif action == 'commit':
            if batch.status != 'reviewed' or request.data.get('acknowledge') is not True:
                raise ValidationError('Review the preview and acknowledge the control totals and source scope.')
            current = make_preview(batch, checked_config(batch.configuration, batch))
            if not current['ready']:
                raise ValidationError({'message': 'Resolve all errors and duplicate decisions; preview again before importing.', 'preview': current})
            if current != batch.preview:
                raise ValidationError('The matching register entries changed since preview. Validate again and review the latest matches.')
            for row in current['rows']:
                if row['excluded']:
                    continue
                validator = EntrySerializer(data=row['payload'])
                validator.is_valid(raise_exception=True)
                entry = validator.save(owner=request.user)
                RegisterImportOrigin.objects.create(batch=batch, entry=entry, file_hash=batch.file_hash, sheet=batch.sheet, source_row=row['row'])
                audit(request.user, 'register.draft_imported', entry.pk, batch=batch.pk, source_row=row['row'])
            batch.preview = current
            batch.status, batch.imported_at = 'imported', timezone.now()
            batch.version += 1
            batch.save()
            audit(request.user, 'register.import_completed', batch.pk, count=current['included'])
        else:
            raise ValidationError('Unsupported import action.')
        return Response(detail(batch))
