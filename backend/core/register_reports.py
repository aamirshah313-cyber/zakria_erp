"""Confirmed register reporting. All amounts and balances remain single-entry movements."""
import calendar
import csv
import io
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape
from django.db import transaction
from django.db.models import Sum, Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response
from rest_framework.views import APIView
from .auth import require, audit
from .models import RegisterEntry, RegisterCategory, RegisterSource, Project, RegisterReportTemplate, Company
from .recovery import desktop_only

FIELDS = {'reporting_class':'Income / expense classification', 'reference':'Register reference', 'date':'Date', 'party':'Organization / person', 'category':'Category', 'project':'Project / contract', 'source':'Cash / bank source', 'direction':'Receipt / payment', 'method':'Payment method', 'instrument':'Instrument / reference', 'handled_by':'Handled by', 'receipt':'Receipts (PKR)', 'payment':'Payments (PKR)', 'balance':'Running balance (PKR)', 'remarks':'Remarks', 'prepared_by':'Prepared by', 'approved_by':'Approved by', 'confirmed_at':'Confirmation timestamp', 'evidence':'Supporting files', 'nature':'Transaction nature', 'beneficiary':'Beneficiary', 'transfer_in':'Transfers in (PKR)', 'transfer_out':'Transfers out (PKR)'}
DEFAULT_COLUMNS = ['date','reference','party','source','receipt','payment','balance']
GROUPS = {'day':'Day', 'quarter':'Fiscal quarter', 'financial_year':'Financial year', 'method':'Payment method', 'nature':'Transaction nature', 'reporting_class':'Income / expense classification', 'beneficiary':'Beneficiary', 'handled_by':'Handled by', 'prepared_by':'Prepared by', 'approved_by':'Approved by', 'month':'Month', 'category':'Category', 'project':'Project / contract', 'source':'Cash / bank source', 'party':'Organization / person'}
SUMMARY_FIELDS = {'label':'Group', 'count':'Statement rows', 'receipt':'Receipts (PKR)', 'payment':'Payments (PKR)', 'transfer_in':'Transfers in (PKR)', 'transfer_out':'Transfers out (PKR)', 'net':'Net movement (PKR)'}
NUMERIC = {'receipt','payment','balance','net','count','evidence','transfer_in','transfer_out'}
BASIS = 'Confirmed receipts less payments within the selected filters. Opening is earlier recorded movement; no independent opening positions are included. Current revision excludes cancelled entries. Net movement is not profit, debt or a reconciled bank balance.'


def parse_day(value):
    if not value:
        return None
    try:
        result = date.fromisoformat(value)
        if not 1900 <= result.year <= 2200:
            raise ValueError()
        return result
    except (TypeError, ValueError):
        raise ValidationError('Select valid calendar dates.')


def date_range(definition):
    period = definition.get('period', 'custom')
    if period == 'custom':
        start, end = parse_day(definition.get('from')), parse_day(definition.get('to'))
    else:
        anchor = parse_day(definition.get('anchor')) or timezone.localdate()
        if period == 'as_at':
            start, end = None, anchor
        elif period == 'day':
            start = end = anchor
        elif period in ['month','quarter','half_year','financial_year']:
            fiscal = definition.get('fiscal_start_month', 7)
            if type(fiscal) is not int or not 1 <= fiscal <= 12:
                raise ValidationError('Select a fiscal-year start month from 1 to 12.')
            months = {'month':1,'quarter':3,'half_year':6,'financial_year':12}[period]
            offset = 0 if months == 1 else (anchor.month - fiscal) % 12 % months
            index = anchor.year * 12 + anchor.month - 1 - offset
            start = date(index // 12, index % 12 + 1, 1)
            last = index + months - 1
            end = date(last // 12, last % 12 + 1, calendar.monthrange(last // 12, last % 12 + 1)[1])
        else:
            raise ValidationError('Select a supported reporting period.')
    if start and end and start > end:
        raise ValidationError('From date cannot be after To date.')
    return start, end


def clean_definition(raw):
    if not isinstance(raw, dict):
        raise ValidationError('Invalid report definition.')
    allowed = ['columns','period','from','to','anchor','fiscal_start_month','category','source','project','source_kind','direction','party','method','group','view','title','header','footer','paper','orientation','accent','chart','counterparty','nature','evidence','balance_basis','reporting_class','beneficiary','handled_by','reference','min_amount','max_amount']
    d = {key: raw[key] for key in allowed if key in raw}
    d.setdefault('columns', DEFAULT_COLUMNS)
    columns = d['columns']
    if not isinstance(columns, list) or not columns or any(type(k) is not str or k not in FIELDS for k in columns) or len(set(columns)) != len(columns):
        raise ValidationError('Select distinct supported report fields.')
    for key, choices, default in [('reporting_class',['','unclassified','income','expense','other'],''), ('balance_basis',['receipts_less_payments','payments_less_receipts'],'receipts_less_payments'), ('evidence',['','missing','present'],''), ('nature',['']+[v[0] for v in RegisterEntry._meta.get_field('nature').choices],''), ('source_kind',['','cash','bank'],''), ('direction',['','receipt','payment'],''), ('method',['','cash','transfer','cheque','card','other'],''), ('group',list(GROUPS),'month'), ('view',['detail','summary'],'detail'), ('paper',['A4','A3'],'A4'), ('orientation',['portrait','landscape'],'landscape')]:
        d.setdefault(key, default)
        if d[key] not in choices:
            raise ValidationError(f'Invalid {key} option.')
    for key, limit, default in [('title',120,'Transaction activity report'),('header',240,''),('footer',160,''),('party',180,''),('beneficiary',180,''),('handled_by',120,''),('reference',120,'')]:
        d.setdefault(key,default)
        if not isinstance(d[key],str) or len(d[key]) > limit or (key == 'title' and not d[key].strip()):
            raise ValidationError(f'{key.title()} must be text up to {limit} characters.')
    d.setdefault('accent','#2563EB')
    if not isinstance(d['accent'],str) or not re.fullmatch(r'#[0-9a-fA-F]{6}',d['accent']):
        raise ValidationError('Use a six-digit report colour such as #2563EB.')
    d.setdefault('chart',True)
    fiscal = d.get('fiscal_start_month', 7)
    if type(fiscal) is not int or not 1 <= fiscal <= 12:
        raise ValidationError('Select a fiscal-year start month from 1 to 12.')
    if type(d['chart']) is not bool:
        raise ValidationError('Chart option must be true or false.')
    for key in ['category','source','project','counterparty']:
        if key in d and d[key] not in [None,'',0]:
            if type(d[key]) is not int or d[key] < 1:
                raise ValidationError('Invalid master-record filter.')
    from .register_imports import exact_amount
    for key in ['min_amount', 'max_amount']:
        if d.get(key) not in [None, '']:
            try:
                d[key] = str(exact_amount(d[key]))
            except ValueError as error:
                raise ValidationError(str(error))
    if d.get('min_amount') and d.get('max_amount') and Decimal(d['min_amount']) > Decimal(d['max_amount']):
        raise ValidationError('Minimum amount cannot exceed maximum amount.')
    date_range(d)
    return d


def report_group(row, d):
    day = date.fromisoformat(row['date'])
    fiscal = d.get('fiscal_start_month', 7)
    if type(fiscal) is not int or not 1 <= fiscal <= 12:
        raise ValidationError('Select a fiscal-year start month from 1 to 12.')
    year = day.year - (day.month < fiscal)
    label = f'FY {year}' if fiscal == 1 else f'FY {year}/{year + 1}'
    if d['group'] == 'quarter':
        return f'{label} Q{((day.month-fiscal) % 12)//3 + 1}'
    if d['group'] == 'financial_year':
        return label
    if d['group'] == 'day':
        return row['date']
    return row['date'][:7] if d['group'] == 'month' else row[d['group']] or 'Unassigned'


def build_report(definition, user_name=''):
    d = clean_definition(definition)
    start, end = date_range(d)
    from .register_movements import movement_rows
    rows, opening, opening_note = movement_rows(d, start, end)
    scope = []
    from .models import Party
    for key, model in [('category',RegisterCategory),('project',Project),('source',RegisterSource),('counterparty',Party)]:
        if d.get(key):
            record = get_object_or_404(model, pk=d[key])
            scope.append(f'{key.title()}: {record.name}')
            if key == 'project':
                scope.extend(filter(None, [f'Contract: {record.reference}' if record.reference else '', f'Client: {record.client.name}' if record.client_id else '', record.location, record.contact_name, record.email]))
    for key in ['source_kind','direction','method','party','nature','evidence','reporting_class','beneficiary','handled_by','reference','min_amount','max_amount']:
        if d.get(key):
            scope.append(f'{key.replace("_", " ").title()}: {d[key]}')
    sign = Decimal(-1) if d['balance_basis'] == 'payments_less_receipts' else Decimal(1)
    opening *= sign
    if not opening:
        opening = Decimal(0)
    running, receipts, payments = opening, Decimal(0), Decimal(0)
    transfer_in, transfer_out = Decimal(0), Decimal(0)
    grouped = defaultdict(lambda: {'receipt':Decimal(0),'payment':Decimal(0),'transfer_in':Decimal(0),'transfer_out':Decimal(0),'count':0})
    for row in rows:
        incoming, outgoing = row['receipt'], row['payment']
        running += sign*(incoming-outgoing+row['transfer_in']-row['transfer_out'])
        receipts += incoming
        payments += outgoing
        transfer_in += row['transfer_in']
        transfer_out += row['transfer_out']
        group = report_group(row, d)
        for key in ['receipt','payment','transfer_in','transfer_out']:
            grouped[group][key] += row[key]
            row[key] = f'{row[key]:.2f}'
        grouped[group]['count'] += 1
        row['balance'] = f'{running:.2f}'
    summary = [{'label':key,'count':v['count'],'receipt':f'{v["receipt"]:.2f}','payment':f'{v["payment"]:.2f}','net':f'{sign*(v["receipt"]-v["payment"]+v["transfer_in"]-v["transfer_out"]):.2f}'} for key,v in sorted(grouped.items())]
    for row in summary:
        for key in ['transfer_in', 'transfer_out']:
            row[key] = f'{grouped[row["label"]][key]:.2f}'
    basis = ('Income and expense figures are explicitly classified cash receipts/payments, not accrual financial statements. Amount filters apply to the original voucher total before allocations. ' if d.get('reporting_class') or d.get('min_amount') or d.get('max_amount') else '') + ('Payments less receipts: positive means net funds applied to this activity. ' if sign == -1 else 'Receipts less payments, plus transfers in less transfers out. ') + opening_note + ' Current revision excludes cancelled records. This is not profit, debt or a reconciled bank balance. Grouped detail can contain allocation rows from the same transaction.'
    company = Company.objects.filter(pk=1).first()
    generated = timezone.localtime().strftime('%Y-%m-%d %H:%M:%S %Z')
    period_label = d.get('period','custom').replace('_',' ').title()
    if d.get('period') in ['quarter','half_year','financial_year']:
        period_label += f' (FY starts {calendar.month_name[d.get("fiscal_start_month",7)]})'
    return {'definition':d,'columns':d['columns'],'rows':rows,'summary':summary,'count':len(rows),'opening':f'{opening:.2f}','receipts':f'{receipts:.2f}','payments':f'{payments:.2f}','closing':f'{running:.2f}','net':f'{running-opening:.2f}','from':str(start or ''),'to':str(end or ''),'period_label':period_label,'scope':'; '.join(scope) or 'All categories, projects, parties and cash/bank sources','basis':basis,'transfers_in':f'{transfer_in:.2f}','transfers_out':f'{transfer_out:.2f}','generated':generated,'prepared_for':user_name,'company':{'name':company.name if company else 'Muhammad Zakaria and Sons','city':company.city if company else 'Islamabad','ntn':company.ntn if company else '', 'ftn':company.ftn if company else '', 'strn':company.strn if company else ''}}


def chart_groups(report):
    rows = sorted(report['summary'],key=lambda r:Decimal(r['receipt'])+Decimal(r['payment']),reverse=True)
    top = rows[:8]
    if len(rows) > 8:
        top.append({'label':'Other groups','receipt':str(sum(Decimal(r['receipt']) for r in rows[8:])),'payment':str(sum(Decimal(r['payment']) for r in rows[8:]))})
    return top


def table_data(report):
    if report['definition']['view'] == 'summary':
        return list(SUMMARY_FIELDS), SUMMARY_FIELDS, report['summary']
    return report['columns'], FIELDS, report['rows']


def metadata(report):
    d = report['definition']
    return [('Company',report['company']['name']),('Report',d['title']),('Period',f'{report["period_label"]}: {report["from"] or "Start of recorded history"} to {report["to"] or "End of recorded history"}'),('Filters',report['scope']),('Generated',report['generated']),('Generated for',report['prepared_for']),('Opening (PKR)',report['opening']),('Receipts (PKR)',report['receipts']),('Payments (PKR)',report['payments']),('Closing (PKR)',report['closing']),('Net movement (PKR)',report['net']),('Transfers in (PKR)',report['transfers_in']),('Transfers out (PKR)',report['transfers_out']),('Basis',report['basis']),('Header',d['header']),('Footer',d['footer'])]


def safe_text(value):
    text = str(value)
    return "'"+text if text.lstrip().startswith(('=','+','-','@')) else text


def csv_bytes(report):
    out = io.StringIO()
    writer = csv.writer(out)
    for key,value in metadata(report):
        writer.writerow([key,safe_text(value)])
    writer.writerow([])
    columns, labels, rows = table_data(report)
    writer.writerow([labels[k] for k in columns])
    for row in rows:
        writer.writerow([str(row[k]) if k in NUMERIC else safe_text(row[k]) for k in columns])
    return ('\ufeff'+out.getvalue()).encode('utf-8')


def xlsx_bytes(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet.pagebreak import Break
    book = Workbook()
    info = book.active
    info.title = 'Report details'
    for key,value in metadata(report):
        info.append([key,safe_text(value)])
    info.column_dimensions['A'].width, info.column_dimensions['B'].width = 25, 110
    for row in info:
        row[1].alignment = Alignment(wrap_text=True,vertical='top')
        info.row_dimensions[row[0].row].height = 44 if len(str(row[1].value)) > 100 else 25
    columns,labels,rows = table_data(report)
    sheet = book.create_sheet('Report rows')
    sheet.append([labels[k] for k in columns])
    for row in rows:
        cells = []
        for key in columns:
            value = row[key]
            if key in NUMERIC:
                value = int(value) if key in ['count','evidence'] else Decimal(value)
            elif key == 'date':
                value = date.fromisoformat(value)
            else:
                value = safe_text(value)
            cells.append(value)
        sheet.append(cells)
    for column,key in enumerate(columns,1):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(column)].width = 20 if key in NUMERIC or key == 'date' else 30
        for cells in sheet.iter_rows(min_row=2,min_col=column,max_col=column):
            cell = cells[0]
            if key in NUMERIC:
                cell.number_format = '#,##0' if key in ['count','evidence'] else '#,##0.00;[Red](#,##0.00)'
            elif key == 'date':
                cell.number_format = 'dd-mmm-yyyy'
            cell.alignment = Alignment(vertical='top',wrap_text=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = '1:1'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = report['definition']['orientation']
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if report['definition']['paper'] == 'A3' else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.print_options.horizontalCentered = True
    sheet.oddHeader.center.text = report['definition']['title'].replace('&','&&')
    sheet.oddFooter.left.text = report['definition']['footer'].replace('&','&&')
    sheet.oddFooter.right.text = 'Page &P of &N'
    summary = book.create_sheet('Grouped summary')
    summary.append(list(SUMMARY_FIELDS.values()))
    for row in report['summary']:
        summary.append([safe_text(row['label']),row['count'],*[Decimal(row[k]) for k in ['receipt','payment','transfer_in','transfer_out','net']]])
    summary.column_dimensions['A'].width = 35
    for key in ['B','C','D','E','F','G']:
        summary.column_dimensions[key].width = 23
    for row in summary.iter_rows(min_row=2,min_col=3,max_col=7):
        for cell in row:
            cell.number_format = '#,##0.00;[Red](#,##0.00)'
    summary.freeze_panes = 'A2'
    if report['summary'] and report['definition']['chart']:
        chart = BarChart()
        chart.title = 'Receipts and payments (PKR)'
        chart.add_data(Reference(summary,min_col=3,max_col=4,min_row=1,max_row=summary.max_row),titles_from_data=True)
        chart.set_categories(Reference(summary,min_col=1,min_row=2,max_row=summary.max_row))
        chart.width,chart.height = 24,12
        chart.series[0].graphicalProperties.solidFill = report['definition']['accent'][1:]
        chart.series[1].graphicalProperties.solidFill = '94A3B8'
        summary.add_chart(chart,'I2')
    accent = report['definition']['accent'][1:]
    brightness = sum(int(accent[i:i+2],16)*w for i,w in [(0,.299),(2,.587),(4,.114)])
    for ws in [sheet,summary]:
        for cell in ws[1]:
            cell.fill = PatternFill('solid',fgColor=accent)
            cell.font = Font(bold=True,color='000000' if brightness > 150 else 'FFFFFF')
            cell.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[1].height = 32
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def pdf_bytes(report):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4,A3,landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.graphics.shapes import Drawing, Rect, String
    d = report['definition']
    size = A3 if d['paper'] == 'A3' else A4
    if d['orientation'] == 'landscape':
        size = landscape(size)
    columns,labels,rows = table_data(report)
    widths = [72 if key in NUMERIC else 65 if key == 'date' else 92 if key in ['reference','method','direction'] else 125 if key != 'remarks' else 170 for key in columns]
    available = size[0]-64
    if sum(widths) > available*1.18:
        raise ValidationError('These columns need more space. Choose A3 landscape or fewer fields for a readable PDF.')
    widths = [w/sum(widths)*available for w in widths]
    if len(rows) > 2000:
        raise ValidationError('PDF detail is limited to 2,000 rows. Narrow the period, choose grouped summary, or export Excel.')
    out = io.BytesIO()
    doc = SimpleDocTemplate(out,pagesize=size,leftMargin=32,rightMargin=32,topMargin=32,bottomMargin=48,title=d['title'],author=report['company']['name'])
    body = ParagraphStyle('Body',fontName='Helvetica',fontSize=8,leading=11,spaceAfter=5,wordWrap='CJK')
    title = ParagraphStyle('Title',parent=body,fontSize=18,leading=22,spaceAfter=12)
    small = ParagraphStyle('Small',parent=body,fontSize=7,leading=10,textColor=colors.HexColor('#475569'))
    accent = colors.HexColor(d['accent'])
    text_colour = colors.black if .299*accent.red+.587*accent.green+.114*accent.blue > .59 else colors.white
    heading = ParagraphStyle('Heading',parent=body,fontName='Helvetica-Bold',textColor=text_colour)
    p = lambda value,style=body: Paragraph(escape(str(value)).replace('\n','<br/>'),style)
    story = [p(report['company']['name'],title),p(d['title'],ParagraphStyle('Subtitle',parent=body,fontSize=12,leading=16))]
    company_ids = ' | '.join(f'{key.upper()}: {report["company"][key]}' for key in ['ntn','ftn','strn'] if report['company'][key])
    if company_ids:
        story.append(p(company_ids,small))
    for key,value in metadata(report):
        if key in ['Period','Filters','Generated','Generated for','Header'] and value:
            story.append(p(f'{key}: {value}',small))
    story += [Spacer(1,8),p(f'Opening PKR {Decimal(report["opening"]):,.2f} | Receipts {Decimal(report["receipts"]):,.2f} | Payments {Decimal(report["payments"]):,.2f} | Closing {Decimal(report["closing"]):,.2f}'),p(report['basis'],small)]
    if Decimal(report['transfers_in']) or Decimal(report['transfers_out']):
        story.append(p(f'Internal transfers: in {Decimal(report["transfers_in"]):,.2f} | out {Decimal(report["transfers_out"]):,.2f}'))
    groups = chart_groups(report)
    if d['chart'] and groups:
        chart = Drawing(available,155)
        peak = max(float(r[key]) for r in groups for key in ['receipt','payment']) or 1
        chart.add(String(0,142,f'Receipts / payments by {GROUPS[d["group"]].lower()} - PKR; top eight groups plus other',fontSize=8))
        step = available/len(groups)
        for i,row in enumerate(groups):
            for j,key in enumerate(['receipt','payment']):
                height = float(row[key])/peak*90
                chart.add(Rect(i*step+8+j*step*.28,30,step*.25,height,fillColor=accent if j == 0 else colors.HexColor('#94A3B8'),strokeColor=None))
            chart.add(String(i*step+8,16,str(row['label'])[:18],fontSize=6))
        chart.add(String(0,2,f'Colour: receipts | Grey: payments | Maximum {peak:,.2f}',fontSize=7))
        story += [chart,Spacer(1,8)]
    def display(key,value):
        return f'{Decimal(str(value)):,.2f}' if key in NUMERIC and key not in ['count','evidence'] else value
    data = [[p(labels[key],heading) for key in columns]] + [[p(display(key,row[key])) for key in columns] for row in rows]
    if not rows:
        story.append(p('No confirmed entries match the selected filters.'))
    table = LongTable(data,colWidths=widths,repeatRows=1,splitByRow=1,splitInRow=1,hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),accent),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F1F5F9')]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
    story.append(table)
    def page(canvas,document):
        canvas.setStrokeColor(accent)
        canvas.line(32,size[1]-18,size[0]-32,size[1]-18)
        canvas.setFont('Helvetica',7)
        footer = p(d['footer'],small)
        _,height = footer.wrap(available-65,30)
        footer.drawOn(canvas,32,15)
        canvas.drawRightString(size[0]-32,20,f'Page {document.page}')
    doc.build(story,onFirstPage=page,onLaterPages=page)
    return out.getvalue()


def image_bytes(report,output):
    from PIL import Image, ImageDraw, ImageFont
    import reportlab
    font_path = Path(reportlab.__file__).parent/'fonts'/'Vera.ttf'
    font = lambda size: ImageFont.truetype(str(font_path),size) if font_path.exists() else ImageFont.load_default(size=size)
    measure = ImageDraw.Draw(Image.new('RGB',(1,1)))
    def wrap(text, size, width=1490):
        lines, line = [], ''
        for word in str(text).split():
            if line and measure.textlength(line+' '+word,font=font(size)) > width:
                lines.append(line)
                line = ''
            for char in word:
                if measure.textlength(line+char,font=font(size)) > width:
                    lines.append(line)
                    line = ''
                line += char
            line += ' '
        if line.strip(): lines.append(line.strip())
        return lines
    header = [(report['company']['name'],25),(report['definition']['title'],28),(f'{report["period_label"]}: {report["from"] or "Start of history"} to {report["to"] or "End of history"}',18),(report['scope'],18),(f'Generated {report["generated"]} for {report["prepared_for"]}',16),(report['definition']['header'],18),('Receipts (report colour) / payments (grey) in PKR',19)]
    heading_lines = [(line,size) for text,size in header for line in wrap(text,size)]
    graph_top = 65+sum(size+10 for _,size in heading_lines)
    groups = chart_groups(report)
    footer_top = graph_top+max(len(groups),1)*69+25
    footers = [('Top eight groups plus Other groups; full group labels are in the tabular report.',17),(f'Total receipts PKR {Decimal(report["receipts"]):,.2f} | Payments {Decimal(report["payments"]):,.2f}',20),(report['basis'],16),(report['definition']['footer'],16)]
    footer_lines = [(line,size) for text,size in footers for line in wrap(text,size)]
    image = Image.new('RGB',(1600,footer_top+sum(size+10 for _,size in footer_lines)+35),'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((0,0,1600,10),fill=report['definition']['accent'])
    y = 35
    for line,size in heading_lines:
        draw.text((55,y),line,font=font(size),fill='#172033')
        y += size+10
    peak = max((float(r[key]) for r in groups for key in ['receipt','payment']),default=1) or 1
    for index,row in enumerate(groups):
        y = graph_top+index*69
        label = row['label']
        while draw.textlength(label,font=font(16)) > 315:
            label = label[:-4]+'...'
        draw.text((55,y),label,font=font(16),fill='#172033')
        for offset,key in [(0,'receipt'),(27,'payment')]:
            width = float(row[key])/peak*930
            draw.rectangle((390,y+offset,390+width,y+offset+20),fill=report['definition']['accent'] if key == 'receipt' else '#94A3B8')
            draw.text((1335,y+offset),f'{Decimal(row[key]):,.2f}',font=font(14),fill='#172033')
    if not groups:
        draw.text((55,graph_top),'No confirmed entries match the selected filters.',font=font(24),fill='#475569')
    y = footer_top
    for line,size in footer_lines:
        draw.text((55,y),line,font=font(size),fill='#475569')
        y += size+10
    buf = io.BytesIO()
    image.save(buf,format='JPEG' if output == 'jpeg' else 'PNG',quality=95,dpi=(150,150))
    return buf.getvalue()


class RegisterReports(APIView):
    def get(self,request):
        desktop_only()
        require(request.user,'register.view')
        return Response({'fields':FIELDS,'default_columns':DEFAULT_COLUMNS,'groups':GROUPS,'today':str(timezone.localdate()),'saved':list(RegisterReportTemplate.objects.filter(owner=request.user,active=True).order_by('name').values('id','name','definition'))})

    def post(self,request):
        desktop_only()
        require(request.user,'register.view')
        output = request.data.get('format','json')
        if output not in ['json','csv','xlsx','pdf','png','jpeg']:
            raise ValidationError('Unsupported report format.')
        if output != 'json':
            require(request.user,'register.export')
        report = build_report(request.data,request.user.username)
        if output == 'json':
            try:
                page = max(1,int(request.data.get('page',1)))
            except (TypeError,ValueError):
                raise ValidationError('Invalid preview page.')
            report['rows'] = report['rows'][(page-1)*100:page*100]
            report['page'] = page
            return Response(report)
        content = {'csv':csv_bytes,'xlsx':xlsx_bytes,'pdf':pdf_bytes}.get(output)
        body = content(report) if content else image_bytes(report,output)
        mime = {'csv':'text/csv; charset=utf-8','xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','pdf':'application/pdf','png':'image/png','jpeg':'image/jpeg'}[output]
        response = HttpResponse(body,content_type=mime)
        response['Content-Disposition'] = f'attachment; filename="register-report.{output}"'
        response['Cache-Control'] = 'no-store'
        audit(request.user,'register.report_exported',output,count=report['count'],view=report['definition']['view'])
        return response


class RegisterReportTemplates(APIView):
    @transaction.atomic
    def post(self,request):
        desktop_only()
        require(request.user,'register.view')
        pk = request.data.get('id')
        item = get_object_or_404(RegisterReportTemplate.objects.select_for_update(),pk=pk,owner=request.user) if pk else None
        if request.data.get('action') == 'archive':
            if not item:
                raise ValidationError('Select a saved report.')
            item.active = False
            item.save()
            audit(request.user,'register.report_archived',item.pk)
            return Response({'message':'Report archived.'})
        name = request.data.get('name','')
        if not isinstance(name,str) or not 1 <= len(name.strip()) <= 100:
            raise ValidationError('Enter a report name of 1–100 characters.')
        definition = clean_definition(request.data.get('definition'))
        if RegisterReportTemplate.objects.filter(owner=request.user,name=name.strip()).exclude(pk=pk).exists():
            raise ValidationError('That saved report name already exists. Choose a different name or update the saved report.')
        if item is None:
            item = RegisterReportTemplate(owner=request.user)
        item.name,item.definition,item.active = name.strip(),definition,True
        item.save()
        audit(request.user,'register.report_saved',item.pk)
        return Response({'id':item.pk,'name':item.name,'definition':item.definition})
